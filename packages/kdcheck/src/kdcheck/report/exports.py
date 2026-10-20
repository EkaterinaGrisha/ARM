"""Human-readable report forms: an Excel workbook and a printable HTML page."""

from __future__ import annotations

import html
import io
from collections.abc import Iterator
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from kdcheck.report import labels as L
from kdcheck.schemas import Extracted, Report, SourceRef, TreeNode


def _v(ex: Extracted | None) -> str | float | int | None:
    return None if ex is None else ex.value


def _walk(nodes: list[TreeNode], level: int = 0) -> Iterator[tuple[int, TreeNode]]:
    for n in nodes:
        yield level, n
        yield from _walk(n.children, level + 1)


def _src(report: Report, src: SourceRef | None) -> str:
    if src is None:
        return ""
    paths = {d.id: d.path for d in report.documents}
    return f"{paths.get(src.document_id, src.document_id)}, стр. {src.page}"


# ------------------------------------------------------------------------ Excel

HEADER_FILL = PatternFill("solid", fgColor="E8EDF5")


def _sheet(
    wb: Workbook, title: str, header: list[str], rows: list[list[Any]], widths: list[int] | None = None
) -> Worksheet:
    ws: Worksheet = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
    for i, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    return ws


def to_xlsx(report: Report) -> bytes:
    wb = Workbook()
    ws = wb.worksheets[0]  # a new workbook always has one worksheet
    ws.title = "Сводка"
    s, k = report.summary, report.kit
    info = [
        ("Комплект", k.name),
        ("Изделие", f"{k.root_designation or '—'} {k.root_name or ''}".strip()),
        ("Обработан", k.processed_at),
        ("Конвейер", f"{k.pipeline_name} {k.pipeline_version}"),
        ("Документов", s.documents_total),
        ("Спецификаций", s.specifications),
        ("Чертежей", s.drawings),
        ("Позиций в дереве", s.tree_items),
        ("Укомплектованность чертежами", f"{s.completeness_ratio:.0%}" if s.completeness_ratio is not None else "—"),
        ("Расхождений", s.discrepancies_total),
        ("Критичных параметров", s.critical_parameters),
    ]
    for key, value in info:
        ws.append([key, value])
    ws.append([])
    ws.append(["Расхождения по типам", ""])
    for t, title in L.DISCREPANCY_TYPES.items():
        ws.append([title, s.discrepancies_by_type.get(t, 0)])
    ws.append([])
    ws.append(["Документы по статусу", ""])
    for st, title in L.DOCUMENT_STATUS.items():
        if s.documents_by_status.get(st):
            ws.append([title, s.documents_by_status[st]])
    ws.column_dimensions["A"].width = 38
    ws.column_dimensions["B"].width = 60
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)

    paths = {d.id: d.path for d in report.documents}
    _sheet(
        wb,
        "Расхождения",
        ["№", "Тип", "Важность", "Обозначение", "Описание", "Ожидалось", "Фактически", "Документы", "Источник"],
        [
            [
                i,
                L.DISCREPANCY_TYPES.get(x.type, x.type),
                L.SEVERITY.get(x.severity, x.severity),
                x.designation,
                x.message,
                x.expected,
                x.actual,
                "\n".join(paths.get(d, d) for d in x.document_ids),
                "\n".join(_src(report, s) for s in x.sources),
            ]
            for i, x in enumerate(report.discrepancies, start=1)
        ],
        [5, 26, 11, 22, 60, 20, 20, 40, 40],
    )

    tree_rows = []
    for level, n in _walk(report.tree):
        tree_rows.append(
            [
                level,
                ("    " * level) + (n.designation or ""),
                n.name,
                L.NODE_KIND.get(n.kind, n.kind),
                n.section,
                n.position,
                n.quantity,
                n.total_quantity,
                L.NODE_STATUS.get(n.status, n.status),
                "\n".join(paths.get(d, d) for d in n.drawing_document_ids),
            ]
        )
    tw = _sheet(
        wb,
        "Дерево изделия",
        [
            "Уровень",
            "Обозначение",
            "Наименование",
            "Вид",
            "Раздел",
            "Поз.",
            "Кол.",
            "Кол. на изделие",
            "Чертёж",
            "Файлы",
        ],
        tree_rows,
        [8, 34, 34, 20, 20, 6, 7, 10, 14, 50],
    )
    for i, (level, _) in enumerate(_walk(report.tree), start=2):
        tw.row_dimensions[i].outlineLevel = min(level, 7)

    _sheet(
        wb,
        "Основные надписи",
        ["Файл"] + list(L.TITLE_FIELDS.values()),
        [[paths.get(p.document_id)] + [_v(getattr(p.title_block, f)) for f in L.TITLE_FIELDS] for p in report.parts],
        [40, 22, 28, 30, 8, 9, 7, 6, 12, 20, 6, 7],
    )

    names = {p.document_id: p.name for p in report.parts}
    _sheet(
        wb,
        "Параметры",
        [
            "Обозначение",
            "Деталь",
            "Категория",
            "Параметр",
            "Значение",
            "Мин.",
            "Макс.",
            "Ед.",
            "Поле допуска / шкала",
            "IT",
            "Общее",
            "Критичный",
            "Исходный текст",
            "Источник",
        ],
        [
            [
                p.designation,
                names.get(p.document_id),
                L.PARAMETER_CATEGORY.get(x.category, x.category),
                x.name,
                x.value,
                x.value_min,
                x.value_max,
                x.unit,
                x.qualifier,
                x.it_grade,
                "да" if x.general else "",
                "да" if x.critical else "",
                x.text,
                _src(report, x.source),
            ]
            for p in report.parts
            for x in p.parameters
        ],
        [20, 22, 22, 22, 9, 8, 8, 6, 14, 5, 7, 9, 50, 40],
    )

    _sheet(
        wb,
        "Критичные параметры",
        ["Обозначение", "Деталь", "Категория", "Параметр", "Правило", "Обоснование", "Оценка", "Источник"],
        [
            [
                c.designation,
                c.part_name,
                L.PARAMETER_CATEGORY.get(c.category, c.category),
                c.text,
                c.rule,
                c.reason,
                c.score,
                _src(report, c.source),
            ]
            for c in report.critical_parameters
        ],
        [20, 22, 20, 24, 20, 70, 8, 40],
    )

    spec_rows = []
    for spec in report.specifications:
        for r in spec.rows:
            spec_rows.append(
                [
                    _v(spec.designation),
                    r.section,
                    _v(r.format),
                    _v(r.position),
                    _v(r.designation),
                    _v(r.name),
                    _v(r.quantity),
                    _v(r.note),
                ]
            )
    _sheet(
        wb,
        "Спецификации",
        ["Спецификация", "Раздел", "Формат", "Поз.", "Обозначение", "Наименование", "Кол.", "Примечание"],
        spec_rows,
        [20, 22, 8, 6, 24, 40, 7, 20],
    )

    _sheet(
        wb,
        "Документы",
        ["Файл", "Тип", "Происхождение", "Стр.", "Обозначение", "Наименование", "Статус", "Проблема", "Примечания"],
        [
            [
                d.path,
                L.DOCUMENT_KIND.get(d.kind, d.kind),
                L.ORIGIN.get(d.origin, d.origin),
                len(d.pages),
                d.designation,
                d.name,
                L.DOCUMENT_STATUS.get(d.status, d.status),
                d.problem,
                "\n".join(d.notes),
            ]
            for d in report.documents
        ],
        [48, 18, 13, 6, 22, 28, 20, 50, 50],
    )

    if report.quality:
        _sheet(
            wb,
            "Качество",
            ["Блок", "Метрика", "Точность (P)", "Полнота (R)", "F1", "Доля верных", "Объём"],
            [[m.block, m.name, m.precision, m.recall, m.f1, m.accuracy, m.support] for m in report.quality.metrics],
            [14, 50, 12, 12, 8, 12, 8],
        )

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------ HTML

CSS = """
body{font:14px/1.45 system-ui,-apple-system,"Segoe UI",Roboto,Arial,sans-serif;color:#1d2433;
margin:32px auto;max-width:1100px;padding:0 16px}
h1{font-size:24px;margin:0 0 4px}h2{font-size:18px;margin:32px 0 8px;border-bottom:1px solid #d8dee9;padding-bottom:4px}
.muted{color:#5b6475}.mono{font-family:ui-monospace,Consolas,monospace}
table{border-collapse:collapse;width:100%;margin:8px 0}
th,td{border:1px solid #d8dee9;padding:4px 8px;text-align:left;vertical-align:top}
th{background:#eef2f8;font-weight:600}.kpi{display:flex;flex-wrap:wrap;gap:12px;margin:16px 0}
.kpi div{border:1px solid #d8dee9;border-radius:8px;padding:8px 14px;min-width:120px}
.kpi b{display:block;font-size:22px}
.high{color:#b42318}.medium{color:#a15c07}.low{color:#475467}ul.tree{list-style:none;padding-left:18px;margin:0}
ul.tree li{margin:1px 0}.missing{color:#b42318;font-weight:600}.warning{color:#a15c07}
@media print{body{margin:0}h2{break-after:avoid}tr{break-inside:avoid}}
"""


def _e(v: object) -> str:
    return "—" if v is None or v == "" else html.escape(str(v))


def _tree_html(nodes: list[TreeNode]) -> str:
    if not nodes:
        return ""
    items = []
    for n in nodes:
        cls = n.status if n.status in ("missing", "warning") else ""
        qty = f" × {n.quantity:g}" if n.quantity else ""
        mark = {"missing": " — нет документа", "warning": " — есть замечания"}.get(n.status, "")
        items.append(
            f'<li class="{cls}"><span class="mono">{_e(n.designation) if n.designation else ""}</span> '
            f"{_e(n.name)}{qty}"
            f'<span class="muted"> · {L.NODE_KIND.get(n.kind, n.kind)}</span>{mark}{_tree_html(n.children)}</li>'
        )
    return f'<ul class="tree">{"".join(items)}</ul>'


def to_html(report: Report) -> str:
    k, s = report.kit, report.summary
    paths = {d.id: d.path for d in report.documents}
    out = [
        "<!doctype html><html lang='ru'><head><meta charset='utf-8'>",
        f"<title>Отчёт — {_e(k.name)}</title><meta name='viewport' content='width=device-width,initial-scale=1'>",
        f"<style>{CSS}</style></head><body>",
        f"<h1>Отчёт о входном контроле КД: {_e(k.name)}</h1>",
        f"<div class='muted'>Изделие: <span class='mono'>{_e(k.root_designation)}</span> {_e(k.root_name)} · "
        f"обработан {_e(k.processed_at)} · конвейер {_e(k.pipeline_name)} {_e(k.pipeline_version)}</div>",
        "<div class='kpi'>",
    ]
    ratio = f"{s.completeness_ratio:.0%}" if s.completeness_ratio is not None else "—"
    for title, value in [
        ("Документов", s.documents_total),
        ("Позиций в дереве", s.tree_items),
        ("Укомплектованность", ratio),
        ("Расхождений", s.discrepancies_total),
        ("Критичных параметров", s.critical_parameters),
    ]:
        out.append(f"<div><b>{_e(value)}</b>{title}</div>")
    out.append("</div>")

    out.append("<h2>Расхождения</h2>")
    if report.discrepancies:
        out.append("<table><tr><th>№</th><th>Тип</th><th>Обозначение</th><th>Описание</th><th>Документы</th></tr>")
        for i, x in enumerate(report.discrepancies, start=1):
            out.append(
                f"<tr><td>{i}</td><td class='{x.severity}'>{_e(L.DISCREPANCY_TYPES.get(x.type, x.type))}</td>"
                f"<td class='mono'>{_e(x.designation)}</td><td>{_e(x.message)}</td>"
                f"<td>{'<br>'.join(_e(paths.get(d, d)) for d in x.document_ids)}</td></tr>"
            )
        out.append("</table>")
    else:
        out.append("<p>Расхождений не найдено.</p>")

    problems = [d for d in report.documents if d.status not in ("ok",)]
    out.append("<h2>Проблемные файлы</h2>")
    if problems:
        out.append("<table><tr><th>Файл</th><th>Статус</th><th>Пояснение</th></tr>")
        for d in problems:
            out.append(
                f"<tr><td>{_e(d.path)}</td><td>{_e(L.DOCUMENT_STATUS.get(d.status, d.status))}</td>"
                f"<td>{_e(d.problem)}</td></tr>"
            )
        out.append("</table>")
    else:
        out.append("<p>Нет.</p>")

    out.append("<h2>Дерево изделия</h2>")
    out.append(_tree_html(report.tree) or "<p>Не построено.</p>")

    out.append("<h2>Основные надписи</h2><table><tr>")
    out.append("".join(f"<th>{t}</th>" for t in list(L.TITLE_FIELDS.values())[:9]) + "</tr>")
    for p in report.parts:
        out.append(
            "<tr>"
            + "".join(f"<td>{_e(_v(getattr(p.title_block, f)))}</td>" for f in list(L.TITLE_FIELDS)[:9])
            + "</tr>"
        )
    out.append("</table>")

    out.append("<h2>Критичные параметры</h2>")
    if report.critical_parameters:
        out.append("<table><tr><th>Обозначение</th><th>Параметр</th><th>Обоснование</th></tr>")
        for c in report.critical_parameters:
            out.append(
                f"<tr><td class='mono'>{_e(c.designation)}</td><td>{_e(c.text)}</td><td>{_e(c.reason)}</td></tr>"
            )
        out.append("</table>")
    else:
        out.append("<p>Нет.</p>")

    out.append("<h2>Параметры технических требований</h2>")
    out.append("<table><tr><th>Обозначение</th><th>Категория</th><th>Параметр</th><th>Исходный текст</th></tr>")
    for p in report.parts:
        for param in p.parameters:
            out.append(
                f"<tr><td class='mono'>{_e(p.designation)}</td><td>{_e(L.PARAMETER_CATEGORY.get(param.category))}</td>"
                f"<td>{_e(param.name)}</td><td>{_e(param.text)}</td></tr>"
            )
    out.append("</table></body></html>")
    return "\n".join(out)
